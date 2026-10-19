"""
Сбор базы заведений и компаний Москвы с телефонами для холодных звонков.

Источник данных: OpenStreetMap через Overpass API (открытая лицензия ODbL).
Результат: Excel-файл с общим листом, листом на каждую категорию и сводкой.

Запуск:
    pip install requests openpyxl
    python parse_moscow.py                    # ~3000 записей -> moscow_leads.xlsx
    python parse_moscow.py --limit 2000 --out leads.xlsx
    python parse_moscow.py --any-site         # не поднимать наверх компании без сайта
"""
import argparse
import json
import re
import sys
import time
from collections import Counter, defaultdict

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OVERPASS_MIRRORS = [
    "https://overpass-api.de/api/interpreter",
    "https://overpass.kumi.systems/api/interpreter",
    "https://maps.mail.ru/osm/tools/overpass/api/interpreter",
]

# Только объекты с телефоном и хотя бы одним «бизнесовым» тегом, в границах Москвы.
QUERY = """
[out:json][timeout:300];
area["name"="Москва"]["admin_level"="4"]->.msk;
(
  nwr(area.msk)["phone"]["name"][~"^(amenity|shop|office|craft|tourism|leisure|healthcare|beauty)$"~"."];
  nwr(area.msk)["contact:phone"]["name"][~"^(amenity|shop|office|craft|tourism|leisure|healthcare|beauty)$"~"."];
);
out center tags;
"""

# (ключ, значение) -> (категория, подкатегория). Значение None — любое.
CATEGORY_RULES = [
    # Общепит
    (("amenity", "restaurant"), ("Рестораны и кафе", "Ресторан")),
    (("amenity", "cafe"), ("Рестораны и кафе", "Кафе")),
    (("amenity", "bar"), ("Рестораны и кафе", "Бар")),
    (("amenity", "pub"), ("Рестораны и кафе", "Паб")),
    (("amenity", "fast_food"), ("Рестораны и кафе", "Фастфуд")),
    (("amenity", "food_court"), ("Рестораны и кафе", "Фуд-корт")),
    (("amenity", "ice_cream"), ("Рестораны и кафе", "Мороженое")),
    (("shop", "bakery"), ("Рестораны и кафе", "Пекарня")),
    (("shop", "confectionery"), ("Рестораны и кафе", "Кондитерская")),
    (("shop", "coffee"), ("Рестораны и кафе", "Кофейня (магазин)")),
    # Красота
    (("shop", "hairdresser"), ("Красота и уход", "Парикмахерская")),
    (("shop", "beauty"), ("Красота и уход", "Салон красоты")),
    (("shop", "cosmetics"), ("Красота и уход", "Косметика")),
    (("shop", "massage"), ("Красота и уход", "Массаж")),
    (("amenity", "spa"), ("Красота и уход", "SPA")),
    (("leisure", "sauna"), ("Красота и уход", "Сауна / баня")),
    (("shop", "tattoo"), ("Красота и уход", "Тату-салон")),
    (("beauty", None), ("Красота и уход", "Салон красоты")),
    # Медицина
    (("amenity", "dentist"), ("Медицина", "Стоматология")),
    (("amenity", "clinic"), ("Медицина", "Клиника")),
    (("amenity", "doctors"), ("Медицина", "Частный врач / центр")),
    (("amenity", "hospital"), ("Медицина", "Больница")),
    (("amenity", "pharmacy"), ("Медицина", "Аптека")),
    (("amenity", "veterinary"), ("Медицина", "Ветклиника")),
    (("shop", "optician"), ("Медицина", "Оптика")),
    (("shop", "medical_supply"), ("Медицина", "Медтехника")),
    (("healthcare", None), ("Медицина", "Медучреждение")),
    # Спорт и досуг
    (("leisure", "fitness_centre"), ("Спорт и фитнес", "Фитнес-клуб")),
    (("leisure", "sports_centre"), ("Спорт и фитнес", "Спортцентр")),
    (("leisure", "dance"), ("Спорт и фитнес", "Танцевальная студия")),
    (("leisure", "swimming_pool"), ("Спорт и фитнес", "Бассейн")),
    (("shop", "sports"), ("Спорт и фитнес", "Спорттовары")),
    (("amenity", "nightclub"), ("Развлечения", "Ночной клуб")),
    (("amenity", "cinema"), ("Развлечения", "Кинотеатр")),
    (("amenity", "theatre"), ("Развлечения", "Театр")),
    (("leisure", "escape_game"), ("Развлечения", "Квест")),
    (("leisure", "bowling_alley"), ("Развлечения", "Боулинг")),
    (("leisure", "amusement_arcade"), ("Развлечения", "Игровой центр")),
    (("amenity", "events_venue"), ("Развлечения", "Площадка для мероприятий")),
    # Гостиницы
    (("tourism", "hotel"), ("Гостиницы", "Отель")),
    (("tourism", "hostel"), ("Гостиницы", "Хостел")),
    (("tourism", "guest_house"), ("Гостиницы", "Гостевой дом")),
    (("tourism", "apartment"), ("Гостиницы", "Апартаменты")),
    (("tourism", "motel"), ("Гостиницы", "Мотель")),
    # Образование
    (("amenity", "language_school"), ("Образование", "Языковая школа")),
    (("amenity", "driving_school"), ("Образование", "Автошкола")),
    (("amenity", "music_school"), ("Образование", "Музыкальная школа")),
    (("amenity", "training"), ("Образование", "Учебный центр")),
    (("amenity", "prep_school"), ("Образование", "Подготовительные курсы")),
    (("amenity", "kindergarten"), ("Образование", "Детский сад")),
    (("amenity", "childcare"), ("Образование", "Детский центр")),
    (("amenity", "college"), ("Образование", "Колледж")),
    # Авто
    (("shop", "car_repair"), ("Автобизнес", "Автосервис")),
    (("shop", "car"), ("Автобизнес", "Автосалон")),
    (("shop", "car_parts"), ("Автобизнес", "Автозапчасти")),
    (("shop", "tyres"), ("Автобизнес", "Шиномонтаж")),
    (("amenity", "car_wash"), ("Автобизнес", "Автомойка")),
    (("amenity", "car_rental"), ("Автобизнес", "Прокат авто")),
    (("craft", "car_repair"), ("Автобизнес", "Автосервис")),
    # Услуги
    (("shop", "dry_cleaning"), ("Бытовые услуги", "Химчистка")),
    (("shop", "laundry"), ("Бытовые услуги", "Прачечная")),
    (("shop", "tailor"), ("Бытовые услуги", "Ателье")),
    (("shop", "shoe_repair"), ("Бытовые услуги", "Ремонт обуви")),
    (("shop", "copyshop"), ("Бытовые услуги", "Копицентр / печать")),
    (("shop", "photo"), ("Бытовые услуги", "Фотоуслуги")),
    (("shop", "mobile_phone"), ("Бытовые услуги", "Салон связи / ремонт")),
    (("shop", "computer"), ("Бытовые услуги", "Компьютеры")),
    (("shop", "electronics"), ("Бытовые услуги", "Электроника")),
    (("shop", "funeral_directors"), ("Бытовые услуги", "Ритуальные услуги")),
    (("shop", "pet_grooming"), ("Бытовые услуги", "Груминг")),
    (("craft", None), ("Бытовые услуги", "Мастерская")),
    # Офисы / B2B
    (("office", "company"), ("Компании и офисы", "Компания")),
    (("office", "estate_agent"), ("Компании и офисы", "Агентство недвижимости")),
    (("office", "lawyer"), ("Компании и офисы", "Юридические услуги")),
    (("office", "notary"), ("Компании и офисы", "Нотариус")),
    (("office", "accountant"), ("Компании и офисы", "Бухгалтерия")),
    (("office", "insurance"), ("Компании и офисы", "Страхование")),
    (("office", "travel_agent"), ("Компании и офисы", "Турагентство")),
    (("shop", "travel_agency"), ("Компании и офисы", "Турагентство")),
    (("office", "employment_agency"), ("Компании и офисы", "Кадровое агентство")),
    (("office", "advertising_agency"), ("Компании и офисы", "Рекламное агентство")),
    (("office", "architect"), ("Компании и офисы", "Архитектурное бюро")),
    (("office", "logistics"), ("Компании и офисы", "Логистика")),
    (("office", "financial"), ("Компании и офисы", "Финансы")),
    (("office", "it"), ("Компании и офисы", "IT-компания")),
    (("office", "coworking"), ("Компании и офисы", "Коворкинг")),
    (("amenity", "coworking_space"), ("Компании и офисы", "Коворкинг")),
    (("office", None), ("Компании и офисы", "Офис")),
    # Магазины
    (("shop", "florist"), ("Магазины", "Цветы")),
    (("shop", "clothes"), ("Магазины", "Одежда")),
    (("shop", "shoes"), ("Магазины", "Обувь")),
    (("shop", "jewelry"), ("Магазины", "Ювелирный")),
    (("shop", "furniture"), ("Магазины", "Мебель")),
    (("shop", "kitchen"), ("Магазины", "Кухни")),
    (("shop", "doityourself"), ("Магазины", "Стройматериалы")),
    (("shop", "hardware"), ("Магазины", "Хозтовары")),
    (("shop", "pet"), ("Магазины", "Зоотовары")),
    (("shop", "toys"), ("Магазины", "Игрушки")),
    (("shop", "books"), ("Магазины", "Книги")),
    (("shop", "gift"), ("Магазины", "Подарки")),
    (("shop", "alcohol"), ("Магазины", "Алкоголь")),
    (("shop", "wine"), ("Магазины", "Вино")),
    (("shop", "butcher"), ("Магазины", "Мясная лавка")),
    (("shop", "supermarket"), ("Магазины", "Супермаркет")),
    (("shop", "convenience"), ("Магазины", "Продукты")),
    (("shop", None), ("Магазины", "Магазин")),
    (("tourism", None), ("Гостиницы", "Туризм")),
    (("leisure", None), ("Развлечения", "Досуг")),
    (("amenity", None), ("Прочее", "Прочее")),
]

# Явно не клиенты: госучреждения, банкоматы и т.п.
EXCLUDE = {
    ("amenity", v) for v in (
        "atm", "bank", "post_office", "police", "townhall", "courthouse",
        "fire_station", "place_of_worship", "school", "university", "library",
        "fuel", "parking", "toilets", "bureau_de_change", "payment_terminal",
        "vending_machine", "parcel_locker", "social_facility", "community_centre",
    )
} | {("office", "government"), ("office", "diplomatic"), ("office", "religion"),
     ("tourism", "attraction"), ("tourism", "museum"), ("tourism", "information"),
     ("leisure", "park"), ("leisure", "playground"), ("leisure", "pitch")}

# Крупные сети — им холодный звонок от агентства почти бесполезен.
CHAIN_KEYWORDS = (
    "пятёрочка", "пятерочка", "магнит", "перекрёсток", "перекресток", "дикси",
    "вкусвилл", "азбука вкуса", "ашан", "лента", "макдоналдс", "kfc", "бургер кинг",
    "вкусно — и точка", "вкусно и точка", "теремок", "шоколадница", "кофе хауз",
    "starbucks", "сбербанк", "сбер", "втб", "мтс", "мегафон", "билайн", "теле2",
    "м.видео", "эльдорадо", "dns", "ригла", "36,6", "горздрав", "столички",
    "озон", "ozon", "wildberries", "яндекс маркет", "лукойл", "роснефть",
)

KEYS = ("amenity", "shop", "office", "craft", "tourism", "leisure", "healthcare", "beauty")
CALL_STATUSES = ["Не звонили", "Недозвон", "Перезвонить", "Не интересно", "Интерес", "Встреча", "Клиент"]


def fetch_overpass(query):
    for url in OVERPASS_MIRRORS:
        for attempt in range(3):
            try:
                print(f"Запрос к {url} (попытка {attempt + 1})...", file=sys.stderr)
                resp = requests.post(url, data={"data": query}, timeout=360,
                                     headers={"User-Agent": "moscow-leads-script/1.0"})
                if resp.status_code == 200:
                    return resp.json()
                print(f"  ответ {resp.status_code}", file=sys.stderr)
            except requests.RequestException as e:
                print(f"  ошибка: {e}", file=sys.stderr)
            time.sleep(5 * (attempt + 1))
    sys.exit("Не удалось получить данные ни с одного Overpass-сервера.")


def normalize_phone(raw):
    """'8 (495) 123-45-67; +7 916 ...' -> ['+7 (495) 123-45-67', '+7 (916) ...']"""
    phones = []
    for part in re.split(r"[;,/]|\s{2,}", raw):
        digits = re.sub(r"\D", "", part)
        if len(digits) == 11 and digits[0] in "78":
            digits = digits[1:]
        elif len(digits) != 10:
            continue
        if digits[0] not in "3489":  # коды РФ: 3xx, 4xx (495/499), 8xx, 9xx
            continue
        phones.append(f"+7 ({digits[:3]}) {digits[3:6]}-{digits[6:8]}-{digits[8:]}")
    return list(dict.fromkeys(phones))


def classify(tags):
    for key, value in ((k, tags[k]) for k in KEYS if k in tags):
        if (key, value) in EXCLUDE:
            return None
    for (key, value), cat in CATEGORY_RULES:
        if key in tags and (value is None or tags[key] == value):
            return cat
    return None


def build_address(tags):
    street = tags.get("addr:street", "")
    house = tags.get("addr:housenumber", "")
    addr = ", ".join(p for p in (street, house) if p)
    return addr or tags.get("addr:full", "")


def parse(data):
    rows, seen_phones = [], set()
    for el in data.get("elements", []):
        tags = el.get("tags", {})
        name = tags.get("name", "").strip()
        cat = classify(tags)
        if not name or not cat:
            continue
        if any(k in name.lower() for k in CHAIN_KEYWORDS) or "brand:wikidata" in tags:
            continue
        phones = normalize_phone(";".join(filter(None, (
            tags.get("phone"), tags.get("contact:phone"), tags.get("contact:mobile")))))
        if not phones or phones[0] in seen_phones:
            continue
        seen_phones.update(phones)
        site = tags.get("website") or tags.get("contact:website") or ""
        osm_type = el["type"]
        rows.append({
            "Название": name,
            "Категория": cat[0],
            "Подкатегория": cat[1],
            "Телефон": phones[0],
            "Доп. телефоны": ", ".join(phones[1:]),
            "Сайт": site,
            "Есть сайт": "Да" if site else "Нет",
            "Email": tags.get("email") or tags.get("contact:email") or "",
            "Адрес": build_address(tags),
            "Часы работы": tags.get("opening_hours", ""),
            "VK / Telegram": tags.get("contact:vk") or tags.get("contact:telegram") or "",
            "Ссылка на карте": f"https://www.openstreetmap.org/{osm_type}/{el['id']}",
            "Источник": "OpenStreetMap",
            "Статус звонка": "Не звонили",
            "Комментарий": "",
        })
    return rows


def select(rows, limit, no_site_first):
    """Равномерно набирает до limit записей по категориям, без сайта — в приоритете."""
    by_cat = defaultdict(list)
    for r in rows:
        by_cat[r["Категория"]].append(r)
    for lst in by_cat.values():
        lst.sort(key=lambda r: (no_site_first and r["Есть сайт"] == "Да", r["Название"].lower()))
    picked = []
    while len(picked) < limit and any(by_cat.values()):
        for cat in sorted(by_cat):
            if by_cat[cat] and len(picked) < limit:
                picked.append(by_cat[cat].pop(0))
    picked.sort(key=lambda r: (r["Категория"], r["Подкатегория"], r["Название"].lower()))
    return picked


HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
NO_SITE_FILL = PatternFill("solid", fgColor="E2EFDA")
WIDTHS = {"Название": 34, "Категория": 20, "Подкатегория": 22, "Телефон": 19, "Доп. телефоны": 22,
          "Сайт": 30, "Есть сайт": 10, "Email": 26, "Адрес": 34, "Часы работы": 24,
          "VK / Telegram": 24, "Ссылка на карте": 22, "Источник": 14, "Статус звонка": 15, "Комментарий": 30}


def write_sheet(ws, rows):
    cols = list(WIDTHS)
    ws.append(cols)
    for c, name in enumerate(cols, 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = WIDTHS[name]
    for r in rows:
        ws.append([r[c] for c in cols])
        if r["Есть сайт"] == "Нет":
            for c in range(1, len(cols) + 1):
                ws.cell(row=ws.max_row, column=c).fill = NO_SITE_FILL
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    status_col = get_column_letter(cols.index("Статус звонка") + 1)
    dv = DataValidation(type="list", formula1='"' + ",".join(CALL_STATUSES) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{status_col}2:{status_col}{max(ws.max_row, 2)}")


def write_excel(rows, path):
    wb = Workbook()
    summary = wb.active
    summary.title = "Сводка"
    summary.append(["Категория", "Всего", "Без сайта"])
    for c in range(1, 4):
        summary.cell(row=1, column=c).font = Font(bold=True, color="FFFFFF")
        summary.cell(row=1, column=c).fill = HEADER_FILL
    total = Counter(r["Категория"] for r in rows)
    no_site = Counter(r["Категория"] for r in rows if r["Есть сайт"] == "Нет")
    for cat in sorted(total):
        summary.append([cat, total[cat], no_site[cat]])
    summary.append(["ИТОГО", sum(total.values()), sum(no_site.values())])
    summary.cell(row=summary.max_row, column=1).font = Font(bold=True)
    summary.append([])
    summary.append(["Зелёным выделены компании без сайта — самые горячие лиды для IT-агентства."])
    summary.append(["Источники: " + ", ".join(sorted({r.get("Источник", "") for r in rows}))
                    + ". Перед звонком сверяйте актуальность номера."])
    summary.column_dimensions["A"].width = 28
    summary.column_dimensions["B"].width = 10
    summary.column_dimensions["C"].width = 12

    write_sheet(wb.create_sheet("Все"), rows)
    for cat in sorted(total):
        write_sheet(wb.create_sheet(cat[:31]), [r for r in rows if r["Категория"] == cat])
    wb.save(path)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--limit", type=int, default=3000, help="сколько записей взять (по умолчанию 3000)")
    ap.add_argument("--out", default="moscow_leads.xlsx", help="путь к Excel-файлу")
    ap.add_argument("--from-json", help="взять сохранённый ответ Overpass вместо запроса")
    ap.add_argument("--save-json", help="сохранить сырой ответ Overpass в файл")
    ap.add_argument("--any-site", dest="no_site_first", action="store_false",
                    help="не поднимать в приоритет компании без сайта")
    args = ap.parse_args()

    if args.from_json:
        with open(args.from_json, encoding="utf-8") as f:
            data = json.load(f)
    else:
        data = fetch_overpass(QUERY)
        if args.save_json:
            with open(args.save_json, "w", encoding="utf-8") as f:
                json.dump(data, f, ensure_ascii=False)

    rows = parse(data)
    print(f"Найдено объектов с телефоном: {len(rows)}", file=sys.stderr)
    rows = select(rows, args.limit, args.no_site_first)
    write_excel(rows, args.out)
    print(f"Сохранено {len(rows)} записей в {args.out}", file=sys.stderr)
    for cat, n in sorted(Counter(r["Категория"] for r in rows).items()):
        print(f"  {cat}: {n}", file=sys.stderr)


if __name__ == "__main__":
    main()
